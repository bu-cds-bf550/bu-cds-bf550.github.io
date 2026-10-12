#!/usr/bin/env python3
"""Assemble a PowerPoint file from decktape's per-slide screenshots.

    tools/pptx_from_screenshots.py <screenshots-dir> <deck-basename> <output.pptx>
    tools/pptx_from_screenshots.py <screenshots-dir> <deck-basename> <output.pptx> --notes <deck.html>

decktape exports the PDF by driving a browser through the deck one slide at a time, and with
`--screenshots` it also saves each slide as a PNG on the way. This script puts those images,
one per slide, into a 16:9 .pptx -- so the PowerPoint is the same rendering as the PDF and the
web deck, pixel for pixel, with none of Pandoc's pptx writer limitations (dropped fragment
pauses, flattened incremental lists, lost theme).

The slides are pictures: the text on them cannot be edited or selected. That is deliberate.
The .pptx exists for students who want to take notes against the slides -- in the notes pane
(every slide gets an empty one) or in text boxes beside the picture -- not as an editable copy
of the deck. The .qmd source is the thing to edit.

With `--notes <deck.html>` the deck's own speaker notes are written into the notes pane
instead, one slide at a time. decktape walks the rendered deck's leaf `<section>` elements in
document order and shoots one screenshot each, so the nth screenshot is the nth leaf section
and the mapping needs no other bookkeeping. The counts are checked before anything is written:
a mismatch means that assumption has broken, and the script stops rather than shifting every
note onto the wrong slide.

decktape names each screenshot <basename>_<index>_<WxH>.png with a 1-based index, and it
joins the screenshots directory with the *full* PDF path, so the files can land in a subtree
of the directory given; the search is recursive for that reason.

Needs python-pptx (slides/requirements.txt). Pillow, which it uses to read image sizes, comes
with matplotlib.
"""
import html
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

INDEX = re.compile(r"_(\d+)_\d+x\d+\.png$")
SECTION = re.compile(r"<section\b[^>]*>|</section>")
NOTES = re.compile(r'<aside class="notes"[^>]*>(.*?)</aside>', re.S)


def screenshots(directory: Path, basename: str) -> list[Path]:
    files = [p for p in directory.rglob(f"{basename}_*_*x*.png") if INDEX.search(p.name)]
    files.sort(key=lambda p: int(INDEX.search(p.name).group(1)))
    return files


def leaf_sections(markup: str) -> list[str]:
    """Every `<section>` in the rendered deck that holds no nested one, in document order.

    Those are reveal.js's slides -- an outer `<section>` is a vertical stack, not a slide --
    and they are what decktape shoots, one screenshot each.
    """
    deck = markup[markup.find('<div class="slides">'):]
    opens: list[int] = []
    spans: list[tuple[int, int]] = []
    for m in SECTION.finditer(deck):
        if m.group(0).startswith("</"):
            if opens:
                spans.append((opens.pop(), m.start()))
        else:
            opens.append(m.start())
    leaves = []
    for start, end in sorted(spans):
        body = deck[start:end]
        if "<section" not in body[body.find(">") + 1:]:
            leaves.append(body)
    return leaves


def note_text(section: str) -> str:
    """The speaker note on one slide, as plain text. Empty when the slide carries none."""
    found = NOTES.search(section)
    if not found:
        return ""
    body = found.group(1)
    body = re.sub(r"</(p|li|div|h\d)>", "\n\n", body)
    body = re.sub(r"<li[^>]*>", "- ", body)
    body = re.sub(r"<br\s*/?>", "\n", body)
    body = re.sub(r"<[^>]+>", "", body)
    body = html.unescape(body)
    body = re.sub(r"[ \t]+", " ", body)
    return re.sub(r"\n{3,}", "\n\n", body).strip()


def speaker_notes(deck_html: Path, expected: int) -> list[str]:
    sections = leaf_sections(deck_html.read_text(encoding="utf-8"))
    if len(sections) != expected:
        raise SystemExit(
            f"pptx_from_screenshots: {deck_html} has {len(sections)} slides but there are "
            f"{expected} screenshots. The notes would land on the wrong slides, so nothing "
            f"was written."
        )
    return [note_text(s) for s in sections]


def build(images: list[Path], output: Path, notes: list[str] | None = None) -> None:
    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 13.333 in: the 16:9 widescreen default
    prs.slide_height = Emu(6858000)   # 7.5 in
    blank = prs.slide_layouts[6]
    for n, image in enumerate(images):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)
        # Created on first access: an empty notes pane, ready for the student -- or, with
        # --notes, the deck's own speaker note for this slide.
        pane = slide.notes_slide
        if notes and notes[n]:
            pane.notes_text_frame.text = notes[n]
    prs.save(str(output))


def main(argv: list[str]) -> int:
    deck_html = None
    if "--notes" in argv:
        at = argv.index("--notes")
        if at + 1 >= len(argv):
            print("pptx_from_screenshots: --notes needs the rendered deck's .html", file=sys.stderr)
            return 2
        deck_html = Path(argv[at + 1])
        argv = argv[:at] + argv[at + 2:]
    if len(argv) != 4:
        print(__doc__.strip().splitlines()[2], file=sys.stderr)
        return 2
    directory, basename, output = Path(argv[1]), argv[2], Path(argv[3])
    images = screenshots(directory, basename)
    if not images:
        print(f"pptx_from_screenshots: no {basename}_*.png under {directory}", file=sys.stderr)
        return 1
    notes = speaker_notes(deck_html, len(images)) if deck_html else None
    build(images, output, notes)
    carried = sum(1 for n in (notes or []) if n)
    print(f"{output}: {len(images)} slides" + (f", {carried} with speaker notes" if notes else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
